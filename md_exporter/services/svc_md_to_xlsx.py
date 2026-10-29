#!/usr/bin/env python3
"""
Markdown to XLSX conversion service
Provides common functionality for converting Markdown tables to XLSX format
"""

import os
import unicodedata
from pathlib import Path
from tempfile import NamedTemporaryFile

import pandas as pd
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from ..utils.markdown_utils import get_md_text
from ..utils.table_utils import SUGGESTED_SHEET_NAME, parse_md_to_tables


def convert_md_to_xlsx(
    md_text: str, output_path: Path, is_strip_wrapper: bool = False, force_text: bool = True
) -> None:
    """
    Convert Markdown tables to XLSX format

    Args:
        md_text: Markdown text to convert
        output_path: Path to save the output XLSX file
        is_strip_wrapper: Whether to remove code block wrapper if present
        force_text: Whether to convert cell values to text type

    Raises:
        ValueError: If input processing or table parsing fails
        Exception: If conversion fails
    """
    # Process Markdown text
    processed_md = get_md_text(md_text, is_strip_wrapper=is_strip_wrapper)

    # Parse Markdown tables
    tables = parse_md_to_tables(processed_md, force_value_to_str=force_text)

    # Convert to XLSX
    with NamedTemporaryFile(suffix=".xlsx", delete=False) as temp_xlsx_file:
        temp_xlsx_path = temp_xlsx_file.name

    try:
        with pd.ExcelWriter(temp_xlsx_path, engine="openpyxl") as writer:
            for i, table in enumerate(tables):
                sheet_name = table.attrs.get(SUGGESTED_SHEET_NAME, f"Sheet{i + 1}")
                table.to_excel(writer, sheet_name=sheet_name, index=False, na_rep="")

                worksheet = writer.sheets[sheet_name]
                # Enable text wrapping for cells that contain line breaks
                for row in worksheet.iter_rows():
                    for cell in row:
                        if cell.value and isinstance(cell.value, str) and "\n" in cell.value:
                            cell.alignment = Alignment(wrap_text=True)

                # Auto-fit column widths (openpyxl worksheets have no autofit())
                _autofit_columns(worksheet)

        # Read temp file and write to target
        output_path.write_bytes(Path(temp_xlsx_path).read_bytes())
    finally:
        # Clean up temporary file
        os.unlink(temp_xlsx_path)


def _get_display_width(text: str) -> int:
    """Approximate display width: wide and full-width characters (e.g. CJK) count as 2, others as 1"""
    return sum(2 if unicodedata.east_asian_width(char) in ("W", "F") else 1 for char in text)


def _autofit_columns(worksheet: Worksheet, max_width: int = 50) -> None:
    """
    Set each column's width to fit its longest line of text, and wrap text in cells
    that are wider than the maximum width so they are not clipped

    Args:
        worksheet: openpyxl worksheet to adjust
        max_width: Maximum width (in characters) a column can grow to, excluding padding
    """
    column_widths: dict[int, int] = {}
    for row in worksheet.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            text_width = max(_get_display_width(line) for line in str(cell.value).split("\n"))
            if text_width > max_width:
                cell.alignment = Alignment(wrap_text=True)
            column_widths[cell.column] = max(column_widths.get(cell.column, 0), text_width)

    for column_index, text_width in column_widths.items():
        worksheet.column_dimensions[get_column_letter(column_index)].width = min(text_width, max_width) + 2
